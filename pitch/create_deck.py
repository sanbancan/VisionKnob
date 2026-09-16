from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "pitch" / "VisionKnob_Pitch_Deck.pptx"
SUNFLOWER = ROOT / "data" / "images" / "sunflower-pacquito-colorado.jpg"
BOULDERS = ROOT / "data" / "images" / "Moeraki-Boulders-New-Zealand.jpg"

NAVY = RGBColor(9, 19, 32)
INK = RGBColor(20, 31, 43)
WHITE = RGBColor(244, 247, 242)
MUTED = RGBColor(171, 188, 197)
CYAN = RGBColor(71, 211, 202)
ORANGE = RGBColor(255, 153, 79)
RED = RGBColor(244, 100, 91)


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def image(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def base(prs, eyebrow, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    text(slide, eyebrow.upper(), 0.55, 0.28, 5.6, 0.25, 9, CYAN, True)
    text(slide, f"{number:02d}", 12.35, 0.28, 0.55, 0.25, 9, MUTED, True, align=PP_ALIGN.RIGHT)
    rect(slide, 0.55, 7.12, 12.2, 0.015, RGBColor(40, 61, 74))
    return slide


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def bullet(slide, value, x, y, color=WHITE, accent=CYAN):
    rect(slide, x, y + 0.13, 0.09, 0.09, accent, True)
    text(slide, value, x + 0.22, y, 5.5, 0.42, 16, color)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 8.15, 0, 5.18, 7.5, INK)
    image(slide, SUNFLOWER, 8.55, 0.55, 3.95, 2.35)
    image(slide, BOULDERS, 8.55, 3.05, 3.95, 3.35)
    text(slide, "VISIONKNOB", 0.72, 0.72, 6.5, 0.5, 15, CYAN, True)
    text(slide, "Ask images\nwithout losing\nyour hands.", 0.7, 1.55, 6.9, 2.35, 34, WHITE, True)
    text(slide, "A tactile local vision assistant powered by Arduino UNO Q,\nModulino controls, Gemma, and GenieX.", 0.75, 4.35, 6.4, 0.85, 16, MUTED)
    rect(slide, 0.75, 5.85, 2.05, 0.43, ORANGE, True)
    text(slide, "LOCAL-FIRST AI", 0.93, 5.9, 1.7, 0.25, 10, NAVY, True, align=PP_ALIGN.CENTER)
    notes(slide, "Open with the core promise: VisionKnob makes image understanding physical and immediate. The user selects an image with a knob and submits a question with a button; the answer is generated locally rather than sent to a cloud service.")

    slide = base(prs, "The friction", 2)
    text(slide, "Image understanding is powerful.\nThe interface is still a bottleneck.", 0.7, 1.0, 7.2, 1.2, 29, WHITE, True)
    bullet(slide, "Finding the right image interrupts the flow.", 0.85, 2.75, MUTED, ORANGE)
    bullet(slide, "Keyboard-first workflows exclude hands-busy moments.", 0.85, 3.45, MUTED, ORANGE)
    bullet(slide, "Cloud-only vision creates privacy and latency concerns.", 0.85, 4.15, MUTED, ORANGE)
    rect(slide, 8.15, 1.45, 3.95, 4.7, RGBColor(18, 39, 50), True)
    text(slide, "SELECT\nASK\nUNDERSTAND", 8.65, 2.05, 3.0, 1.75, 25, CYAN, True, align=PP_ALIGN.CENTER)
    text(slide, "Three actions.\nOne physical loop.", 8.8, 4.55, 2.65, 0.7, 17, WHITE, True, align=PP_ALIGN.CENTER)
    notes(slide, "Frame the problem as interaction friction, not a lack of model capability. Vision models can answer questions, but the path from a physical image library to a grounded answer is still awkward.")

    slide = base(prs, "The solution", 3)
    text(slide, "A physical control surface\nfor a local vision model.", 0.7, 0.95, 6.8, 1.2, 30, WHITE, True)
    rect(slide, 0.75, 2.65, 3.35, 2.55, RGBColor(18, 39, 50), True)
    text(slide, "01", 1.05, 2.95, 0.7, 0.5, 23, ORANGE, True)
    text(slide, "Turn", 1.05, 3.6, 2.6, 0.45, 23, WHITE, True)
    text(slide, "Choose the image", 1.05, 4.2, 2.5, 0.35, 15, MUTED)
    rect(slide, 4.95, 2.65, 3.35, 2.55, RGBColor(18, 39, 50), True)
    text(slide, "02", 5.25, 2.95, 0.7, 0.5, 23, CYAN, True)
    text(slide, "Press A", 5.25, 3.6, 2.6, 0.45, 23, WHITE, True)
    text(slide, "Submit the typed question", 5.25, 4.2, 2.8, 0.35, 15, MUTED)
    rect(slide, 9.15, 2.65, 3.35, 2.55, RGBColor(18, 39, 50), True)
    text(slide, "03", 9.45, 2.95, 0.7, 0.5, 23, ORANGE, True)
    text(slide, "See", 9.45, 3.6, 2.6, 0.45, 23, WHITE, True)
    text(slide, "Get a grounded answer", 9.45, 4.2, 2.8, 0.35, 15, MUTED)
    notes(slide, "Explain the interaction in three beats. The browser remains useful for typing and seeing the answer, while the Arduino controller removes the need to reach for the mouse or keyboard during selection and submission.")

    slide = base(prs, "Under the hood", 4)
    text(slide, "One loop. Four layers.", 0.7, 0.95, 6.5, 0.6, 30, WHITE, True)
    layers = [("MODULINO", "Knob + buttons", ORANGE), ("UNO Q", "USB serial bridge", CYAN), ("STREAMLIT", "Selection + chat", ORANGE), ("GEMMA", "Grounded vision", CYAN)]
    for index, (title, subtitle, color) in enumerate(layers):
        x = 0.85 + index * 3.05
        rect(slide, x, 2.45, 2.35, 1.65, RGBColor(18, 39, 50), True)
        text(slide, title, x + 0.2, 2.78, 1.95, 0.3, 13, color, True, align=PP_ALIGN.CENTER)
        text(slide, subtitle, x + 0.2, 3.35, 1.95, 0.35, 14, WHITE, True, align=PP_ALIGN.CENTER)
        if index < 3:
            text(slide, ">", x + 2.48, 3.0, 0.35, 0.45, 23, MUTED, True, align=PP_ALIGN.CENTER)
    text(slide, "JSON events in. Status commands out.\nThe image stays local; the answer stays grounded in the selected frame.", 1.0, 5.15, 11.0, 0.75, 18, MUTED, align=PP_ALIGN.CENTER)
    notes(slide, "Walk left to right. The hardware emits JSON events over USB serial. Streamlit owns the interaction state and attaches the selected image to the request. GenieX serves the cached Gemma VLM locally. Status commands return to the Pixels module for immediate feedback.")

    slide = base(prs, "The demo", 5)
    text(slide, "A question becomes a visible event.", 0.7, 0.95, 8.5, 0.6, 29, WHITE, True)
    image(slide, SUNFLOWER, 0.8, 2.0, 4.25, 3.4)
    rect(slide, 5.55, 2.0, 6.6, 3.4, RGBColor(18, 39, 50), True)
    text(slide, "TURN KNOB", 6.0, 2.45, 2.0, 0.3, 12, ORANGE, True)
    text(slide, "Select the sunflower image", 6.0, 2.85, 4.7, 0.4, 20, WHITE, True)
    text(slide, "TYPE QUESTION", 6.0, 3.55, 2.0, 0.3, 12, CYAN, True)
    text(slide, "What colors are visible?", 6.0, 3.95, 4.7, 0.4, 20, WHITE, True)
    text(slide, "PRESS A", 6.0, 4.65, 2.0, 0.3, 12, ORANGE, True)
    text(slide, "Gemma answers from this image", 6.0, 5.05, 4.8, 0.4, 20, WHITE, True)
    notes(slide, "Use this as the live demo slide. Demonstrate the knob changing the selected image, type a short question, press Button A, and point to the Pixel status feedback while the local model processes the request.")

    slide = base(prs, "Why local", 6)
    text(slide, "Private by default.\nResponsive by design.", 0.7, 0.95, 7.2, 1.15, 30, WHITE, True)
    for index, (heading, body, color) in enumerate([
        ("PRIVATE", "Images remain on the\nlocal machine.", CYAN),
        ("GROUNDED", "The selected frame is\nattached to each request.", ORANGE),
        ("TACTILE", "The controller gives\nfeedback at a glance.", CYAN),
    ]):
        x = 0.85 + index * 4.05
        rect(slide, x, 2.75, 3.35, 2.25, RGBColor(18, 39, 50), True)
        text(slide, heading, x + 0.3, 3.15, 2.75, 0.35, 14, color, True, align=PP_ALIGN.CENTER)
        text(slide, body, x + 0.35, 3.75, 2.65, 0.75, 19, WHITE, True, align=PP_ALIGN.CENTER)
    notes(slide, "The important differentiator is not merely that the model is local. The full interaction is local: image storage, retrieval, inference, and physical feedback all happen without sending the image to a hosted vision API.")

    slide = base(prs, "The hardware", 7)
    text(slide, "Small modules. Clear roles.", 0.7, 0.95, 7.2, 0.6, 30, WHITE, True)
    image(slide, BOULDERS, 8.8, 1.35, 3.5, 4.9)
    components = [("KNOB", "Select image", ORANGE), ("BUTTON A", "Submit question", CYAN), ("BUTTON B", "Next image", ORANGE), ("PIXELS", "Ready / working / result", CYAN)]
    for index, (heading, body, color) in enumerate(components):
        y = 2.0 + index * 0.92
        rect(slide, 0.9, y, 0.12, 0.5, color)
        text(slide, heading, 1.25, y - 0.02, 2.0, 0.28, 13, color, True)
        text(slide, body, 3.25, y - 0.02, 4.65, 0.32, 17, WHITE, True)
    notes(slide, "Tie each physical component to a user outcome. Button C is intentionally open for a future action such as clearing the conversation, toggling retrieval mode, or returning to the first image.")

    slide = base(prs, "The ask", 8)
    text(slide, "Make image understanding\nfeel as immediate as looking.", 0.7, 1.0, 7.8, 1.35, 30, WHITE, True)
    rect(slide, 0.75, 3.1, 5.25, 1.75, CYAN, True)
    text(slide, "NEXT", 1.15, 3.43, 1.0, 0.25, 12, NAVY, True)
    text(slide, "Button C + richer\nphysical feedback", 1.15, 3.85, 3.9, 0.65, 22, NAVY, True)
    rect(slide, 6.55, 3.1, 5.25, 1.75, ORANGE, True)
    text(slide, "NOW", 6.95, 3.43, 1.0, 0.25, 12, NAVY, True)
    text(slide, "Try the local\nvision loop", 6.95, 3.85, 3.9, 0.65, 22, NAVY, True)
    text(slide, "VisionKnob  |  Arduino UNO Q  |  Gemma  |  GenieX", 0.75, 6.18, 8.8, 0.3, 13, MUTED, True)
    notes(slide, "Close with the invitation: try the loop. The prototype already proves the interaction. The next step is expanding the control vocabulary, especially assigning Button C to a high-value workflow action and refining feedback for hands-busy use cases.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()